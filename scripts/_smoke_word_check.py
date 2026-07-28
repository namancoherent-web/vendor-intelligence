"""Ultra-small pipeline run to inspect Word (company + functionality) output.

Target: finish in well under 15 min. Writes files under output/demo/.
"""
from __future__ import annotations

import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

# Shrink discovery/validation BEFORE settings load — aim ~10 min
os.environ["MAX_VALIDATION_ENTITIES"] = "12"
os.environ["TARGET_UNIQUE_COMPANIES"] = "20"
os.environ["TARGET_SOLID_COMPANIES"] = "12"
os.environ["VOLUME_PROMPT_COUNT"] = "6"
os.environ["PHASE1_SMOKE_MAX_PROMPTS"] = "3"
os.environ["PHASE1_GLOBAL_SMOKE_MAX_PROMPTS"] = "4"
os.environ["PHASE3_AGENTIC_VALIDATION"] = "false"
os.environ["PHASE3_AGENTIC_MAX_ENTITIES"] = "0"
os.environ["PHASE3_AGENTIC_MAX_LLM_CALLS"] = "0"
os.environ["PHASE3_PARALLEL_WORKERS"] = "4"
os.environ["DDG_WORKER_COUNT"] = "2"

from vendor_intel.placeholders.load_keys import apply_env_overrides

apply_env_overrides()

from run_query import run_one_query, _parse_input
from ui.bootstrap import load_settings, pipeline_caps
from vendor_intel.pipeline.cap_profiles import apply_cap

MARKET = "bamboo toothbrushes market"
OUT_NAME = "smoke_bamboo_toothbrushes"


def main() -> None:
    settings = load_settings("quality")
    settings = apply_cap(settings, "focused")
    # Ultra-tight for a ~10 min end-to-end check
    settings = settings.model_copy(
        update={
            "pipeline_discover_max": 25,
            "pipeline_global_discover_max": 25,
            "pipeline_enrich_max": 25,
            "pipeline_global_enrich_max": 25,
            "pipeline_export_max_rows": 15,
            "pipeline_global_export_max_rows": 15,
            "volume_prompt_count": 6,
            "pipeline_global_volume_prompt_count": 6,
            "widen_loop_max": 1,
            "max_validation_entities": 12,
        }
    )
    market, country, sections = _parse_input(MARKET)
    classify_cap, enrich_cap = pipeline_caps(settings, country=country)
    print(f"START market={market!r} country={country} classify={classify_cap} enrich={enrich_cap}", flush=True)
    t0 = time.time()
    result = run_one_query(
        market,
        country,
        settings,
        enrich_limit=enrich_cap,
        classify_limit=classify_cap,
        sections=sections,
        out_name=OUT_NAME,
    )
    elapsed = time.time() - t0
    rows = [r for r in (result.get("relevant_companies") or []) if r.get("is_relevant")]
    print(f"DONE in {elapsed/60:.1f} min — {len(rows)} relevant companies", flush=True)
    print(f"  csv : {result.get('_csv_path')}", flush=True)
    print(f"  xlsx: {result.get('_xlsx_path')}", flush=True)
    print(f"  docx: {result.get('_docx_path')}", flush=True)
    print(f"  json: {result.get('_json_path')}", flush=True)

    # Inspect Word: company name vs functionality per line
    docx_path = result.get("_docx_path") or ""
    if not docx_path or not Path(docx_path).exists():
        print("WORD MISSING — docx was not written (check DOCX export failed above)", flush=True)
        return

    from docx import Document

    doc = Document(docx_path)
    print("\n=== WORD DOCUMENT CONTENTS ===", flush=True)
    blank_company = 0
    blank_func = 0
    listed = 0
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        # Numbered company lines look like: "1. Acme Corp (does X)"
        if text[:1].isdigit() and ". " in text[:6]:
            listed += 1
            # split "N. Name (func)" 
            after_num = text.split(". ", 1)[1]
            if " (" in after_num and after_num.endswith(")"):
                name, _, rest = after_num.partition(" (")
                func = rest[:-1] if rest.endswith(")") else rest
            else:
                name, func = after_num, ""
            name = name.strip()
            func = func.strip()
            if not name:
                blank_company += 1
                print(f"  [NO NAME] {text!r}", flush=True)
            elif not func:
                blank_func += 1
                print(f"  [NO FUNC] {name!r}", flush=True)
            else:
                print(f"  OK  {name!r}  |  {func[:80]!r}", flush=True)
        else:
            print(f"  :: {text[:100]}", flush=True)

    print(
        f"\nSUMMARY: {listed} company lines, "
        f"{blank_company} missing name, {blank_func} missing functionality",
        flush=True,
    )


if __name__ == "__main__":
    main()
